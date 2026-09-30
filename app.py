import streamlit as st
import os, io, time, json
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import google.generativeai as genai

# Configuration de la page Streamlit
st.set_page_config(
    page_title="SN LE JOKER FICHE - Pédagogie CEB Sénégal",
    page_icon="🇸🇳",
    layout="wide",
    initial_sidebar_state="expanded"
)

# -------------------------------------------------------------
# SYSTEME DE SÉCURITÉ / MOT DE PASSE
# -------------------------------------------------------------
MOT_DE_PASSE_EXIGE = "Le joker 10"

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

def check_password():
    if st.session_state.get("password_input") == MOT_DE_PASSE_EXIGE:
        st.session_state.authenticated = True
        del st.session_state["password_input"]  # Supprime le mot de passe de la mémoire par sécurité
    else:
        st.session_state.authenticated = False
        st.error("🔒 Mot de passe incorrect. Veuillez réessayer.")

# Écran de verrouillage si l'utilisateur n'est pas connecté
if not st.session_state.authenticated:
    st.title("🔒 Accès Sécurisé")
    st.info("Veuillez saisir le mot de passe pour accéder à SN LE JOKER FICHE avant de pouvoir générer des fiches.")
    
    st.text_input(
        "Mot de passe :", 
        type="password", 
        key="password_input", 
        on_change=check_password
    )
    st.button("Se connecter", on_click=check_password, type="primary")
    st.stop()  # Arrête le chargement du reste de la page tant que le mdp n'est pas bon

# -------------------------------------------------------------
# CONFIGURATION GEMINI API
# -------------------------------------------------------------
api_key = os.environ.get("GEMINI_API_KEY")

if api_key:
    genai.configure(api_key=api_key)

# -------------------------------------------------------------
# CHARGEMENT DE LA BASE DE DONNÉES LOCALE (fiches_ce2.json)
# -------------------------------------------------------------
DB_FILE = 'fiches_ce2.json'
fiches_db = {}

if os.path.exists(DB_FILE):
    try:
        with open(DB_FILE, 'r', encoding='utf-8') as f:
            data_loaded = json.load(f)
            if isinstance(data_loaded, list):
                for item in data_loaded:
                    key = f"{item.get('classe')}_{item.get('domaine')}_{item.get('sous_domaine', '')}_{item.get('notion')}"
                    fiches_db[key] = item
            elif isinstance(data_loaded, dict):
                fiches_db = data_loaded
    except Exception as e:
        st.error(f"Erreur de lecture de la base locale : {e}")

# Style CSS
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;600;700;800&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }
    
    .sn-banner {
        background: linear-gradient(135deg, #0b6623 0%, #15803d 40%, #047857 100%);
        border-radius: 16px;
        padding: 24px;
        color: white;
        margin-bottom: 24px;
        box-shadow: 0 10px 25px -5px rgba(11, 102, 35, 0.25);
        border: 1px solid rgba(255, 255, 255, 0.15);
        position: relative;
        overflow: hidden;
    }
    .joker-pill {
        background: #ffcc00;
        color: #0f172a;
        font-size: 0.8rem;
        font-weight: 800;
        letter-spacing: 1.5px;
        padding: 5px 14px;
        border-radius: 9999px;
        display: inline-flex;
        align-items: center;
        gap: 6px;
        box-shadow: 0 2px 6px rgba(0,0,0,0.15);
        text-transform: uppercase;
    }
    .app-title {
        font-size: 2.1rem;
        font-weight: 800;
        margin: 10px 0 4px 0;
        color: #ffffff;
        letter-spacing: -0.5px;
    }
    .app-subtitle {
        font-size: 0.95rem;
        color: #e2e8f0;
        margin: 0;
    }
    .fiche-box {
        background: #ffffff;
        border-radius: 12px;
        border: 2px solid #0b6623;
        padding: 30px;
        box-shadow: 0 8px 30px rgba(0,0,0,0.06);
        color: #1e293b;
    }
    
    div[data-testid="stSidebar"] button[kind="primary"] {
        background-color: #e31b23 !important;
        background: linear-gradient(135deg, #e31b23 0%, #b91c1c 100%) !important;
        color: white !important;
        font-weight: 800 !important;
        font-size: 1.05rem !important;
        border-radius: 10px !important;
        border: 2px solid #ffcc00 !important;
        padding: 14px 20px !important;
        box-shadow: 0 6px 18px rgba(227, 27, 35, 0.45) !important;
        text-transform: uppercase !important;
        letter-spacing: 0.5px !important;
    }
</style>
""", unsafe_allow_html=True)

# BANDEAU D'EN-TÊTE
st.markdown("""
<div class="sn-banner">
    <span class="joker-pill">🃏 SIGNATURE : LE JOKER</span>
    <h1 class="app-title">SN LE JOKER FICHE (HYBRIDE)</h1>
    <p class="app-subtitle">
        Générateur Intelligent de Fiches Pédagogiques CEB ultra-détaillées et contextualisées au Sénégal
    </p>
</div>
""", unsafe_allow_html=True)

# ARBORESCENCE CEB
DISCIPLINES_TREE = {
    "Français": [
        "Grammaire", "Conjugaison", "Orthographe", "Vocabulaire", "Communication Orale", "Lecture / Langage", "Production d'Écrits (Rédaction)"
    ],
    "Mathématiques": [
        "Calcul", "Activités Numériques (Nombres & Opérations)", "Calcul Mental",
        "Activités Géométriques", "Activités de Mesure", "Résolution de Problèmes"
    ],
    "Éducation à la Science et à la Vie Sociale (ESVS)": [
        "Découverte du Monde - Histoire du Sénégal", "Découverte du Monde - Géographie du Sénégal",
        "Découverte du Monde - Initiation Scientifique & Tech. (IST)",
        "Développement Durable - Vivre ensemble (Civisme & Morale)",
        "Développement Durable - Vivre dans son milieu (Hygiène, Santé, Environnement)"
    ],
    "Éducation Artistique & EPS": [
        "Arts Plastiques (Dessin & Modelage)", "Éducation Musicale (Chant & Rythmes)", "Éducation Physique et Sportive (EPS)"
    ]
}

CLASSES_LIST = [
    ("CI", "Étape 1", "CI (Cours d'Initiation - Étape 1)"),
    ("CP", "Étape 1", "CP (Cours Préparatoire - Étape 1)"),
    ("CE1", "Étape 2", "CE1 (Cours Élémentaire 1 - Étape 2)"),
    ("CE2", "Étape 2", "CE2 (Cours Élémentaire 2 - Étape 2)"),
    ("CM1", "Étape 3", "CM1 (Cours Moyen 1 - Étape 3)"),
    ("CM2", "Étape 3", "CM2 (Cours Moyen 2 - Étape 3)")
]

if "fiche_data" not in st.session_state:
    st.session_state.fiche_data = None

# SIDEBAR
with st.sidebar:
    st.markdown("### 🇸🇳 Profil & Établissement")
    prof_nom = st.text_input("Initiales / Visa enseignant", "Le joker")
    ia_nom = st.text_input("IA", "IA Dakar")
    ief_nom = st.text_input("IEF", "IEF Pikine - Guédiawaye")
    ecole_nom = st.text_input("École", "École Élémentaire Publique Dalifort")
    
    st.markdown("---")
    col_g, col_f = st.columns(2)
    with col_g:
        n_garcons = st.number_input("Garçons (G)", 0, 100, 25)
    with col_f:
        n_filles = st.number_input("Filles (F)", 0, 100, 23)
    effectif_str = f"{n_garcons + n_filles} élèves (G: {n_garcons} | F: {n_filles})"
    date_fiche = st.date_input("Date", date.today())
    
    st.markdown("---")
    
    # FORMULAIRE DE SELECTION
    sel_classe_label = st.selectbox("Classe", [c[2] for c in CLASSES_LIST], index=3)
    classe_code = sel_classe_label.split(" ")[0]
    cycle_code = "Étape 1" if classe_code in ["CI", "CP"] else ("Étape 2" if classe_code in ["CE1", "CE2"] else "Étape 3")

    sel_discipline = st.selectbox("Discipline", list(DISCIPLINES_TREE.keys()), index=0)
    sel_activite = st.selectbox("Sous-domaine / Activité", DISCIPLINES_TREE[sel_discipline], index=0)
    notion_input = st.text_input("Notion / Titre de la leçon", "Les noms propres")

    st.markdown("<br>", unsafe_allow_html=True)

# BOUTON DE GÉNÉRATION
if "last_generation_time" not in st.session_state:
    st.session_state.last_generation_time = 0

COOLDOWN_SECONDS = 10
elapsed_time = time.time() - st.session_state.last_generation_time

cle_locale = f"{classe_code}_{sel_discipline}_{sel_activite}_{notion_input}"

if cle_locale in fiches_db or elapsed_time >= COOLDOWN_SECONDS:
    if st.sidebar.button("🃏 GÉNÉRER LE JOKER FICHE", type="primary"):
        if cle_locale in fiches_db:
            st.session_state.fiche_data = fiches_db[cle_locale]
            st.toast("⚡ Fiche chargée depuis la base locale (0 API consommée)", icon="🚀")
        else:
            if not api_key:
                st.error("⚠️ La clé API Gemini est manquante. Configurez GEMINI_API_KEY dans les Secrets Streamlit.")
            else:
                st.session_state.last_generation_time = time.time()
                with st.spinner("🃏 L'IA Le Joker rédige la fiche pédagogique CEB..."):
                    prompt = f"""
Tu es un expert mondial en ingénierie pédagogique selon l'Approche Par Compétences (APC) au Sénégal.
Rédige une fiche pédagogique CEB ULTRA-DÉTAILLÉE ET EXPLICITE :
- Classe : {classe_code} ({cycle_code})
- Discipline : {sel_discipline}
- Activité : {sel_activite}
- Notion : {notion_input}

CONSIGNES STRICTES :
- NE DONNE AUCUNE CONSIGNE GÉNÉRIQUE OU VAGUE.
- Sois précis, explicite, et ancré dans le contexte sénégalais.

Réponds EXCLUSIVEMENT sous la forme d'un objet JSON valide structuré ainsi :
{{
  "titre": "{notion_input}",
  "classe": "{classe_code}",
  "domaine": "{sel_discipline}",
  "sous_domaine": "{sel_activite}",
  "cb": "Texte de la Compétence de Base adaptée",
  "os": "Texte de l'Objectif Spécifique",
  "materiel": "Matériel et supports didactiques",
  "duree": "30 min",
  "etapes": [
    {{
      "nom": "Révision / Pré-requis",
      "duree": "5 min",
      "activite_maitre": "Consignes et questions précises de l'enseignant",
      "activite_eleve": "Réponses attendues des élèves"
    }},
    {{
      "nom": "Situation d'apprentissage (Observation / Découverte)",
      "duree": "10 min",
      "activite_maitre": "Présentation du texte ou support et questions de découverte",
      "activite_eleve": "Lecture, observation et compréhension"
    }},
    {{
      "nom": "Analyse / Synthèse",
      "duree": "10 min",
      "activite_maitre": "Questions de guidage et conceptualisation",
      "activite_eleve": "Dégagement de la règle ou du savoir-faire"
    }},
    {{
      "nom": "Évaluation / Transfert",
      "duree": "5 min",
      "activite_maitre": "Exercice ou consigne d'évaluation précise",
      "activite_eleve": "Résolution individuelle"
    }}
  ]
}}
"""
                    try:
                        model = genai.GenerativeModel('gemini-2.5-flash')
                        response = model.generate_content(prompt)

                        res_text = response.text.strip()
                        if res_text.startswith("```json"):
                            res_text = res_text[7:-3].strip()
                        elif res_text.startswith("```"):
                            res_text = res_text[3:-3].strip()

                        data = json.loads(res_text)
                        st.session_state.fiche_data = data
                        st.success("🎉 Fiche générée avec succès via l'IA !")

                    except Exception as e:
                        st.error(f"Erreur lors de la génération : {str(e)}")
else:
    remaining = int(COOLDOWN_SECONDS - elapsed_time)
    st.sidebar.warning(f"⏳ Veuillez patienter {remaining} s avant une génération IA.")
    st.sidebar.button("🃏 GÉNÉRER LE JOKER FICHE", disabled=True)

# FONCTION D'EXPORT WORD (.DOCX)
def generate_docx(data):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sn = p_top.add_run("RÉPUBLIQUE DU SÉNÉGAL\nMINISTÈRE DE L'ÉDUCATION NATIONALE\n")
    r_sn.font.bold = True
    r_sn.font.size = Pt(11)
    r_sn.font.color.rgb = RGBColor(11, 102, 35)
    
    r_title = p_top.add_run("FICHE PÉDAGOGIQUE DE PRÉPARATION (CEB)\n")
    r_title.font.bold = True
    r_title.font.size = Pt(13)
    
    table_admin = doc.add_table(rows=3, cols=2)
    table_admin.style = 'Table Grid'
    c = table_admin.rows
    c[0].cells[0].text = f"IA : {ia_nom} | IEF : {ief_nom}\nÉcole : {ecole_nom}"
    c[0].cells[1].text = f"Classe : {classe_code} ({cycle_code})\nEffectif : {effectif_str}"
    c[1].cells[0].text = f"Discipline : {sel_discipline} - {sel_activite}\nNotion : {notion_input}"
    c[1].cells[1].text = f"Durée : {data.get('duree', '30 min')}\nDate : {date_fiche.strftime('%d/%m/%Y')}"
    c[2].cells[0].text = f"Matériel : {data.get('materiel', 'Tableau, cahiers, manuels')}"
    c[2].cells[1].text = f"Visa enseignant : {prof_nom}"
    
    doc.add_paragraph()
    p_cadre = doc.add_paragraph()
    p_cadre.add_run("I. CADRAGE PÉDAGOGIQUE\n").font.bold = True
    p_cadre.add_run(f"• CB : {data.get('cb', '')}\n")
    p_cadre.add_run(f"• OS : {data.get('os', '')}\n").font.bold = True
    
    # Tableau de déroulement
    p_seq = doc.add_paragraph()
    p_seq.add_run("\nII. DÉROULEMENT DIDACTIQUE DE LA SÉANCE\n").font.bold = True
    
    table_deroul = doc.add_table(rows=1, cols=4)
    table_deroul.style = 'Table Grid'
    headers = ["PHASES / ÉTAPES", "ACTIVITÉS DU MAÎTRE", "ACTIVITÉS DES ÉLÈVES", "DURÉE"]
    for i, h in enumerate(headers):
        cell = table_deroul.rows[0].cells[i]
        cell.text = h
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0B6623"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    etapes = data.get('etapes', [])
    for etape in etapes:
        row_cells = table_deroul.add_row().cells
        row_cells[0].text = etape.get('nom', '')
        row_cells[1].text = etape.get('activite_maitre', '')
        row_cells[2].text = etape.get('activite_eleve', '')
        row_cells[3].text = etape.get('duree', '')
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# AFFICHAGE DE LA FICHE GÉNÉRÉE
if st.session_state.fiche_data:
    d = st.session_state.fiche_data
    
    st.markdown("### 👁️ Aperçu de la Fiche Pédagogique")
    
    # Bouton de Téléchargement Word
    docx_bytes = generate_docx(d)
    st.download_button(
        label="📥 Télécharger la Fiche Officielle Word (.docx)",
        data=docx_bytes,
        file_name=f"Fiche_CEB_{classe_code}_{sel_activite}_{prof_nom}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )
    
    st.markdown(f"""
    <div class="fiche-box" style="margin-top:15px;">
        <h3 style="color:#0b6623; text-align:center;">{sel_discipline} - {sel_activite}</h3>
        <h4>📌 Notion : {notion_input}</h4>
        <hr>
        <p><b>• Compétence de Base (CB) :</b> {d.get('cb')}</p>
        <p><b>• Objectif Spécifique (OS) :</b> {d.get('os')}</p>
        <p><b>• Matériel didactique :</b> {d.get('materiel')}</p>
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("#### ⏱️ Déroulement Didactique Détaillé")
    
    etapes = d.get('etapes', [])
    if etapes:
        import pandas as pd
        df_etapes = pd.DataFrame(etapes)
        df_etapes = df_etapes.rename(columns={
            'nom': 'ÉTAPES / PHASES',
            'activite_maitre': 'ACTIVITÉS DU MAÎTRE',
            'activite_eleve': 'ACTIVITÉS DES ÉLÈVES',
            'duree': 'DURÉE'
        })
        st.dataframe(df_etapes, use_container_width=True, hide_index=True)

else:
    st.info("👈 Veuillez sélectionner vos options dans le panneau de gauche et cliquer sur **« 🃏 GÉNÉRER LE JOKER FICHE »**.")